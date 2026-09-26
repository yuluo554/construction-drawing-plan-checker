# -*- coding: utf-8 -*-
"""施工方案解析器（docx）—— 规则版 v0（不依赖 LLM）。

流程：python-docx 读取 → 标题树（Heading样式/编号正则双通道）→ 表格统计 →
     正则参数抽取（混凝土/脚手架/基坑/高支模/工程概况），每值带章节号+原文证据。
LLM 结构化抽取接口 extract_with_llm() 留待 v2（需 API Key）。
"""
import os
import re
from typing import List

from ..schemas.parameter_card import ParameterCard, Parameter, Evidence

# ---------- 章节树 ----------

NUM_HEADING_RE = re.compile(r"^\s*((?:第[一二三四五六七八九十百\d]+[章节部分])|(?:\d+(?:\.\d+)+))[、\s]*(\S.*)$")
CN_NUM_HEADING_RE = re.compile(r"^\s*([一二三四五六七八九十]+)[、.]\s*(\S.*)$")
MAX_LEVEL_LEN = 40


def build_outline(doc) -> List[dict]:
    """返回章节树（扁平列表，含 level/编号/标题/段落序号）。"""
    outline = []
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if not text:
            continue
        level = None
        style = ""
        try:
            style = para.style.name or ""
        except Exception:
            pass
        m = NUM_HEADING_RE.match(text)
        mc = CN_NUM_HEADING_RE.match(text)
        if "Heading" in style or "标题" in style:
            try:
                level = int(re.search(r"(\d+)", style).group(1))
            except Exception:
                level = 1
            num = ""
            if m and len(text) <= MAX_LEVEL_LEN:
                num = m.group(1)
            elif mc and len(text) <= MAX_LEVEL_LEN:
                num = mc.group(1)
            outline.append({"level": level, "num": num, "title": text, "pidx": i})
            continue
        elif m and len(text) <= MAX_LEVEL_LEN:
            num = m.group(1)
            level = 1 if ("第" in num or "章" in num) else num.count(".") + 1
            outline.append({"level": level, "num": num, "title": m.group(2), "pidx": i})
            continue
        elif mc and len(text) <= MAX_LEVEL_LEN:
            outline.append({"level": 1, "num": mc.group(1), "title": mc.group(2), "pidx": i})
            continue
        if level is not None and len(text) <= MAX_LEVEL_LEN:
            outline.append({"level": level, "num": "", "title": text, "pidx": i})
    return outline


def _current_section(outline, pidx):
    """段落所属最小章节号（如 '3.2' 或 '四'）。"""
    sec = ""
    for item in outline:
        if item["pidx"] <= pidx and item["num"]:
            sec = item["num"]
    return sec


def _last_heading(outline, pidx):
    title = ""
    for item in outline:
        if item["pidx"] <= pidx:
            title = item["title"]
    return title


# ---------- 正则抽取规则 ----------

def _num(text):
    try:
        return float(text)
    except (TypeError, ValueError):
        return None


M_RE = r"(\d+(?:\.\d+)?)\s*(?:m|米)(?![a-zA-Z])"

# (entity, attr, pattern, unit)  —— pattern 需含一个捕获组=数值
PARAM_RULES = [
    ("脚手架工程", "立杆纵距", re.compile(r"立杆[纵]距[为约是：:\s]*" + M_RE), "m"),
    ("脚手架工程", "立杆横距", re.compile(r"立杆[横]距[为约是：:\s]*" + M_RE), "m"),
    ("脚手架工程", "步距", re.compile(r"(?:立杆)?步距[为约是：:\s]*" + M_RE), "m"),
    ("脚手架工程", "搭设高度", re.compile(r"(?:脚手架)?搭设高度[为约是：:\s]*" + M_RE), "m"),
    ("基坑工程", "基坑开挖深度", re.compile(r"(?:基坑|基槽|土方开挖)(?:的开挖)?深度[为约是：:\s]*" + M_RE), "m"),
    ("模板支撑体系", "支撑搭设高度", re.compile(r"(?:支撑|支模|模板支架)高度[为约是：:\s]*" + M_RE), "m"),
    ("模板支撑体系", "支撑搭设高度", re.compile(r"搭设高度[为约是：:\s]*" + M_RE + r"[^\n]{0,20}(?:支(?:撑|模)架|高支模)"), "m"),
    ("模板支撑体系", "立杆间距", re.compile(r"(?:立杆|支撑架立杆)(?:纵|横)?间距[为约是：:\s]*" + M_RE), "m"),
]
CONC_RULE = re.compile(r"(?:混凝土强度等级|混凝土)[为：:\s]*(C\d{2,3})")
CONC_PART = re.compile(r"(基础底板|底板|基础|筏板|承台|地梁|墙|柱|梁|板|楼梯|构造柱|圈梁|屋面)")
OVERVIEW_RULES = [
    ("建筑面积", re.compile(r"(?:总建筑面积|建筑面积)[为约是：:\s]*([\d,.]+)\s*(?:㎡|平方米|m2|m²|万㎡)")),
    ("建筑高度", re.compile(r"建筑高度[为约是：:\s]*" + M_RE)),
    ("结构形式", re.compile(r"结构(?:形式|类型)\s*[为约是：:]\s*([\u4e00-\u9fa5A-Za-z]{2,20})")),
    ("基础形式", re.compile(r"基础(?:形式|类型)\s*[为约是：:]\s*([\u4e00-\u9fa5A-Za-z]{2,20})")),
    ("建筑层数", re.compile(r"(?:建筑)?层数\s*[为约是：:]\s*([地上地下\d\-F/B＋+]{1,15})")),
]

# 表格"键单元格|值单元格"模式：(键正则, 类型, 默认实体/属性)
KV_HEIGHT_KEY = re.compile(r"(?:模板支架|支撑架|支模架|模板支撑|支架)[^，。\n]{0,4}高度")
KV_DIST_KEY = re.compile(r"立杆[纵横]?间距|立杆[纵横]距|(?:最大|顶层|水平杆)?步距")
KV_CONC_KEY = re.compile(r"混凝土强度等级?")
NUM_CELL_RE = re.compile(r"^\d+(?:\.\d+)?$")


def _row_kv_params(cells, ev, row_text):
    """从一行单元格中提取 键→值 参数（单位统一换算为 m）。"""
    out = []
    part = CONC_PART.search(row_text)
    for i, text in enumerate(cells):
        if not text:
            continue
        if i + 1 >= len(cells):
            continue
        val_text = cells[i + 1].strip()
        mnum = NUM_CELL_RE.match(val_text)
        if KV_CONC_KEY.search(text):
            mg = re.fullmatch(r"(C\d{2,3})", val_text)
            if mg:
                entity = "混凝土" + ("(%s)" % part.group(1) if part else "")
                out.append(_mk(entity, "强度等级", mg.group(1), "", ev, category="material"))
            continue
        if not mnum:
            continue
        val = float(val_text)
        if KV_HEIGHT_KEY.search(text):
            height = val if re.search(r"\(\s*m\s*\)", text) else val / 1000.0
            if 1 <= height <= 50:
                out.append(_mk("模板支撑体系", "支撑搭设高度", ("%g" % height), "m", ev))
        elif KV_DIST_KEY.search(text):
            dist = val if re.search(r"\(\s*m\s*\)", text) else val / 1000.0
            if 0.3 <= dist <= 20:
                attr = "步距" if "步距" in text else "立杆间距"
                out.append(_mk("模板支撑体系" if "支架" in row_text or "支撑" in row_text
                               else "脚手架工程", attr, ("%g" % dist), "m", ev))
    return out


def _mk(entity, attr, value, unit, ev, source="plan", category="measure"):
    return Parameter(entity=entity, category=category, attr=attr, value=value,
                     unit=unit, conf=0.9, source=source, evidence=ev)


def extract_parameters(doc, outline, filename) -> List[Parameter]:
    params = []
    for i, para in enumerate(doc.paragraphs):
        text = para.text.strip()
        if not text:
            continue
        sec = _current_section(outline, i)
        ev = Evidence(doc=filename, page=i, section=sec, raw=text[:160])

        for entity, attr, pat, unit in PARAM_RULES:
            m = pat.search(text)
            if m:
                v = _num(m.group(1))
                if v is not None and 0.1 <= v <= 200:
                    params.append(_mk(entity, attr, m.group(1), unit, ev))

        for mc in CONC_RULE.finditer(text):
            part = CONC_PART.search(text)
            entity = "混凝土" + ("(%s)" % part.group(1) if part else "")
            params.append(_mk(entity, "强度等级", mc.group(1), "", ev, category="material"))

        for attr, pat in OVERVIEW_RULES:
            m = pat.search(text)
            if m:
                params.append(_mk("工程概况", attr, m.group(1).strip("，。 "), "", ev,
                                  category="overview"))
    return params


def _iter_body(document):
    """按正文顺序产出 ('p', 段落对象) / ('tbl', 表格对象, 其前段落计数)。"""
    from docx.table import Table
    from docx.text.paragraph import Paragraph
    pcount = 0
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            yield ("p", Paragraph(child, document), pcount)
            pcount += 1
        elif child.tag.endswith("}tbl"):
            yield ("tbl", Table(child, document), pcount)


def extract_from_tables(document, outline, filename) -> List[Parameter]:
    """扫描表格单元格中的参数（混凝土强度/脚手架参数常在表格中）。"""
    params = []
    for kind, obj, pcount in _iter_body(document):
        if kind != "tbl":
            continue
        sec = _current_section(outline, max(pcount - 1, 0))
        try:
            rows = obj.rows
        except Exception:
            continue
        seen_texts = set()
        for row in rows:
            cell_texts = []
            for cell in row.cells:
                text = cell.text.strip()
                cell_texts.append(text)
                if not text or text in seen_texts:
                    continue
                seen_texts.add(text)
                ev = Evidence(doc=filename, page=-1, section=sec, raw=text[:160])
                for entity, attr, pat, unit in PARAM_RULES:
                    m = pat.search(text)
                    if m:
                        v = _num(m.group(1))
                        if v is not None and 0.1 <= v <= 200:
                            params.append(_mk(entity, attr, m.group(1), unit, ev))
                for mc in CONC_RULE.finditer(text):
                    part = CONC_PART.search(text)
                    entity = "混凝土" + ("(%s)" % part.group(1) if part else "")
                    params.append(_mk(entity, "强度等级", mc.group(1), "", ev, category="material"))
                for attr, pat in OVERVIEW_RULES:
                    m = pat.search(text)
                    if m:
                        params.append(_mk("工程概况", attr, m.group(1).strip("，。 "), "", ev,
                                          category="overview"))
            row_ev = Evidence(doc=filename, page=-1, section=sec,
                              raw=" | ".join(c for c in cell_texts if c)[:160])
            row_text = " ".join(cell_texts)
            params.extend(_row_kv_params(cell_texts, row_ev, row_text))
    return params


# ---------- 主入口 ----------

def parse_docx(path: str, use_llm: bool = False) -> ParameterCard:
    card = ParameterCard(doc_type="plan", file=os.path.abspath(path))
    try:
        import docx as docx_lib
        document = docx_lib.Document(path)
    except Exception as exc:
        card.warnings.append("docx 打开失败: %s" % exc)
        return card

    filename = os.path.basename(path)
    outline = build_outline(document)
    card.extras["outline"] = [{"level": o["level"], "num": o["num"], "title": o["title"]}
                              for o in outline[:200]]
    card.extras["table_count"] = len(document.tables)
    card.extras["para_count"] = len(document.paragraphs)

    card.meta["标题(首段)"] = next((p.text.strip() for p in document.paragraphs if p.text.strip()), "")
    card.meta["章节树深度"] = max([o["level"] for o in outline], default=0)

    for p in extract_parameters(document, outline, filename):
        card.add(p)
    for p in extract_from_tables(document, outline, filename):
        card.add(p)

    if use_llm:
        from ..llm.extract_plan import extract_with_llm
        llm_params, stats = extract_with_llm(document, outline, filename)
        card.extras["llm_stats"] = stats
        # 合并：LLM 结果去重后追加（规则抽取的确定性优先）
        seen = {(p.entity, p.attr, p.value) for p in card.parameters}
        for p in llm_params:
            if p.entity == "__error__":
                card.warnings.append(p.value)
                continue
            if (p.entity, p.attr, p.value) not in seen:
                seen.add((p.entity, p.attr, p.value))
                card.add(p)

    if not outline:
        card.warnings.append("未识别出章节树（可能为扫描版或非常规格式）")
    return card


def extract_with_llm(path):
    """LLM 结构化抽取接口（v2 接入大模型；v0 返回空）。"""
    return []
