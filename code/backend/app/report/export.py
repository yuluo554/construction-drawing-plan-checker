# -*- coding: utf-8 -*-
"""校核报告 docx 导出 —— 汇总表 + 分级问题清单（含双侧证据/依据/建议）+ 签署栏。"""
import os
import time


def build_report_docx(report, drawing_card=None, plan_card=None, out_path="校核报告.docx"):
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    doc.add_heading("施工图与施工方案协同校核报告", level=0)

    # 项目信息
    meta = (plan_card.meta if plan_card else {}) or {}
    doc.add_paragraph("生成时间：%s" % time.strftime("%Y-%m-%d %H:%M:%S"))
    if plan_card is not None:
        doc.add_paragraph("方案文件：%s" % os.path.basename(plan_card.file or "-"))
    if drawing_card is not None:
        dn = drawing_card.meta.get("project_name") or os.path.basename(drawing_card.file or "-")
        doc.add_paragraph("图纸文件：%s" % dn)

    # 结论与汇总
    doc.add_heading("一、校核结论", level=1)
    p = doc.add_paragraph()
    run = p.add_run(str(report.get("结论", "")))
    run.bold = True
    s = report.get("summary", {})
    color = {"致命": RGBColor(0xC0, 0x00, 0x00), "一般": RGBColor(0xBF, 0x8F, 0x00),
             "提示": RGBColor(0x00, 0x70, 0xC0)}
    doc.add_paragraph("致命：%d    一般：%d    提示：%d" % (s.get("致命", 0), s.get("一般", 0), s.get("提示", 0)))

    # 问题清单
    doc.add_heading("二、问题清单", level=1)
    findings = report.get("findings", [])
    if findings:
        tbl = doc.add_table(rows=1, cols=6)
        tbl.style = "Table Grid"
        headers = ["级别", "规则", "问题", "图纸侧", "方案侧", "建议"]
        for i, h in enumerate(headers):
            c = tbl.cell(0, i)
            c.text = h
            c.paragraphs[0].runs[0].bold = True
        for f in findings:
            row = tbl.add_row().cells
            row[0].text = str(f.get("级别", ""))
            row[1].text = str(f.get("rule_id", ""))
            row[2].text = "%s（%s）" % (f.get("问题", ""), f.get("类型", ""))
            ev = f.get("证据") or {}
            row[3].text = str(f.get("图纸", ev.get("图纸", "")))[:60]
            row[4].text = str(f.get("方案", ev.get("方案", "")))[:60]
            row[5].text = str(f.get("建议", ""))[:60]
            if f.get("级别") in color and row[0].paragraphs[0].runs:
                row[0].paragraphs[0].runs[0].font.color.rgb = color[f["级别"]]
                row[0].paragraphs[0].runs[0].bold = True
            if f.get("条文"):
                row[1].text += "\n依据：%s" % f["条文"].get("path", "")[:30]
    else:
        doc.add_paragraph("未发现问题线索。")

    # 解析过程
    doc.add_heading("三、处理过程", level=1)
    for t in report.get("_trace", []):
        doc.add_paragraph("节点 %s：%s（%.2fs）%s" % (
            t.get("node"), "成功" if t.get("ok") else "失败", t.get("耗时s", 0),
            t.get("error", "")))

    # 签署
    doc.add_heading("四、签署", level=1)
    doc.add_paragraph("校核：____________    复核：____________    批准：____________")
    doc.add_paragraph("（本报告由智能校核系统生成，问题线索需人工确认后处置）")

    doc.save(out_path)
    return out_path
